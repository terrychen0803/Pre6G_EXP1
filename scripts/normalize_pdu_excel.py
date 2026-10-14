"""Normalize a PDU Excel export into explicit timezone-aware watt samples."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path


def offset(value: str) -> timezone:
    match = re.fullmatch(r"([+-])(\d{2}):(\d{2})", value)
    if not match or int(match[2]) > 23 or int(match[3]) > 59:
        raise ValueError("--timezone-offset must be like +08:00")
    minutes = int(match[2]) * 60 + int(match[3])
    return timezone(timedelta(minutes=minutes * (1 if match[1] == "+" else -1)))


def timestamp(value: object, tz: timezone) -> str:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
    else:
        raise ValueError(f"unsupported PDU timestamp: {value!r}")
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        parsed = parsed.replace(tzinfo=tz)
    return parsed.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def normalize(input_path: Path, output_path: Path, *, sheet: str | None, header_row: int,
              timestamp_column: str, power_column: str, outlet_column: str | None,
              outlet_value: str | None, timezone_offset: str, unit: str) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("install requirements-evaluation.txt to read Excel exports") from exc
    if header_row < 1 or unit not in {"W", "kW"}:
        raise ValueError("invalid header row or power unit")
    if bool(outlet_column) != bool(outlet_value):
        raise ValueError("outlet column and outlet value must be provided together")
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.active
        sheet_name = worksheet.title
        values = worksheet.iter_rows(values_only=True)
        headers = None
        for index, row in enumerate(values, start=1):
            if index == header_row:
                headers = [str(value).strip() if value is not None else "" for value in row]
                break
        if headers is None or len(set(headers)) != len(headers):
            raise ValueError("header row is missing or contains duplicate column names")
        required = [timestamp_column, power_column] + ([outlet_column] if outlet_column else [])
        if any(name not in headers for name in required):
            raise ValueError(f"Excel sheet lacks required columns: {required}")
        indexes = {name: headers.index(name) for name in required}
        tz = offset(timezone_offset)
        rows = []
        for row in values:
            if all(value is None for value in row):
                continue
            if outlet_column and str(row[indexes[outlet_column]]).strip() != outlet_value:
                continue
            raw_power = row[indexes[power_column]]
            watts = float(raw_power) * (1000 if unit == "kW" else 1)
            if not math.isfinite(watts) or watts < 0:
                raise ValueError(f"invalid PDU power value: {raw_power!r}")
            rows.append((timestamp(row[indexes[timestamp_column]], tz), watts))
    finally:
        workbook.close()
    if len(rows) < 2:
        raise ValueError("at least two PDU samples are required after outlet filtering")
    rows.sort(key=lambda item: item[0])
    instants = [datetime.fromisoformat(item[0].replace("Z", "+00:00")) for item in rows]
    if any(right <= left for left, right in zip(instants, instants[1:])):
        raise ValueError("PDU timestamps are duplicated or unordered")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        raise ValueError(f"refusing to overwrite {output_path}")
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["timestamp", "power_w"])
        writer.writerows(rows)
    manifest = {"schema_version": "pre6g.pdu-excel-normalization/v1", "source_file": str(input_path.resolve()),
                "source_sha256": hashlib.sha256(input_path.read_bytes()).hexdigest(), "sheet": sheet_name,
                "header_row": header_row, "timestamp_column": timestamp_column, "power_column": power_column,
                "outlet_column": outlet_column, "outlet_value": outlet_value, "timezone_offset_for_naive_values": timezone_offset,
                "source_power_unit": unit, "sample_count": len(rows), "first_timestamp": rows[0][0],
                "last_timestamp": rows[-1][0], "normalized_csv": str(output_path.resolve())}
    output_path.with_suffix(".provenance.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert a PDU .xlsx export to timestamp,power_w CSV without assuming its sampling semantics")
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--timestamp-column", required=True)
    parser.add_argument("--power-column", required=True)
    parser.add_argument("--outlet-column")
    parser.add_argument("--outlet-value")
    parser.add_argument("--timezone-offset", required=True, help="Offset for Excel timestamps without a timezone, e.g. +08:00")
    parser.add_argument("--unit", choices=("W", "kW"), required=True)
    args = parser.parse_args()
    print(json.dumps(normalize(args.input, args.output, sheet=args.sheet, header_row=args.header_row,
                               timestamp_column=args.timestamp_column, power_column=args.power_column,
                               outlet_column=args.outlet_column, outlet_value=args.outlet_value,
                               timezone_offset=args.timezone_offset, unit=args.unit), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
